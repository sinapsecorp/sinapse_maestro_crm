from sqlalchemy.orm import Session
from datetime import datetime, date
from src.database.session import SessionLocal
from src.services.email_service import email_service
from src.services.marketing_service import marketing_service
from src.utils.crypto import decrypt_str
from src.repositories.lead_repository import lead_repository
from src.repositories.campaign_repository import campaign_repository
from src.utils.queue import get_redis_connection
from src.repositories.area_of_expertise_repository import area_of_expertise_repository
from src.models.campaign_lead import CampaignLead, EmailStatus
import uuid
import os
from src.models.campaign import Campaign, CampaignStatus
from src.models.lead import Lead
from email_validator import validate_email, EmailNotValidError
import os
import csv
from io import StringIO
import openpyxl
import xlrd
import unicodedata


def import_leads_from_rows(rows: list[dict], user_id: str, default_area_id: str | None = None, *, progress_base: int = 0, progress_total: int | None = None) -> int:
    db: Session = SessionLocal()
    created = 0
    updated = 0
    skipped = 0
    try:
        total = len(rows)
        # Atualiza meta inicial caso ainda não tenha sido setado
        try:
            from rq import get_current_job
            job = get_current_job()
            if job and not job.meta:
                job.meta = {"progress": 0, "inserted": 0, "updated": 0, "skipped": 0}
                job.save_meta()
        except Exception:
            pass
        for idx, row in enumerate(rows, start=1):
            # Normalizar chaves vindas de XLS onde cabeçalhos podem variar
            def g(*keys):
                for k in keys:
                    v = row.get(k)
                    if v not in (None, ""):
                        return v
                return None

            # Permitir que a área seja nome (string) além de UUID
            area_value = g("area_of_expertise_id", "area", "AreaId", "Area")
            # Se vier um nome de área, buscar o id correspondente
            if area_value and isinstance(area_value, str) and len(area_value) > 20 and '-' in area_value:
                area_id = area_value
            else:
                # tentar por nome
                area = None
                if isinstance(area_value, str) and area_value.strip():
                    area = area_of_expertise_repository.get_by_name(db, name=area_value.strip())
                area_id = str(area.id) if area else area_value

            # telefones: mapear segundo layout e normalizar um principal
            raw_phone = g("phone", "Telefone", "telefone", "Phone", "Telefone Principal", "telefone principal")
            raw_phone2 = g("Telefone Secundário", "telefone secundario", "telefone secund  e1rio")
            def norm_phone(v):
                if v in (None, ""): return None
                digits = ''.join([c for c in str(v) if c.isdigit()])
                if not digits: return None
                if digits.startswith('55') and len(digits) in (12,13):
                    return "+" + digits
                return "+55" + digits
            phone_norm = norm_phone(raw_phone) or norm_phone(raw_phone2)

            company = g("company", "Empresa", "empresa", "Estabelecimentos", "estabelecimentos")
            nome = g("full_name", "name", "nome", "Nome", "Razão Social", "razao social", "Raz e3o Social")

            cnpj = g("CNPJ", "cnpj")
            razao_social = g("Raz e3o Social", "Razao Social", "razao social", "RAZAO SOCIAL")
            if not razao_social:
                razao_social = nome

            data = {
                # obrigatórios do layout
                "cnpj": cnpj,
                "razao_social": razao_social,

                # dados mapeados
                "porte_codigo": g("Porte C f3digo", "Porte Codigo", "Porte C f3digo"),
                "porte": g("Porte"),
                "capital_social": g("Capital Social"),
                "natureza_juridica_codigo": g("Natureza Jur eddica C f3digo", "Natureza Juridica Codigo"),
                "natureza_juridica": g("Natureza Jur eddica", "Natureza Juridica"),
                "ente_federativo_responsavel": g("Ente Federativo Respons elvel", "Ente Federativo Responsavel"),
                "tipo": g("Tipo"),
                "data_abertura": g("Data de Abertura"),
                "nome_fantasia": g("Nome Fantasia"),
                "situacao_cadastral_codigo": g("Situa e7 e3o Cadastral C f3digo", "Situacao Cadastral Codigo"),
                "situacao_cadastral": g("Situa e7 e3o Cadastral", "Situacao Cadastral"),
                "situacao_cadastral_data": g("Situa e7 e3o Cadastral Data", "Situacao Cadastral Data"),
                "situacao_motivo_codigo": g("Situa e7 e3o Motivo C f3digo", "Situacao Motivo Codigo"),
                "situacao_motivo": g("Situa e7 e3o Motivo", "Situacao Motivo"),
                "situacao_especial_codigo": g("Situa e7 e3o Especial C f3digo", "Situacao Especial Codigo"),
                "situacao_especial": g("Situa e7 e3o Especial", "Situacao Especial"),
                "situacao_especial_data": g("Situa e7 e3o Especial Data", "Situacao Especial Data"),
                "telefone_principal": raw_phone,
                "telefone_secundario": raw_phone2,
                "phone": phone_norm,
                "email": g("email", "Email", "E-mail", "e-mail"),
                "municipio_ibge": g("Munic edpio IBGE", "Municipio IBGE"),
                "logradouro": g("Logradouro"),
                "numero": g("N famero", "Numero"),
                "complemento": g("Complemento"),
                "bairro": g("Bairro"),
                "cidade": g("Cidade"),
                "estado": g("Estado"),
                "cep": g("CEP"),
                "pais_codigo": g("Pais C f3digo", "Pais Codigo"),
                "pais": g("Pais", "Pais"),
                "atividade_principal_codigo": g("Atividade Principal C f3digo", "Atividade Principal Codigo"),
                "atividade_principal": g("Atividade Principal"),
                "ultima_atualizacao": g(" daltima Atualiza e7 e3o", "Ultima Atualizacao"),

                # compatibilidade com telas atuais
                "full_name": nome,
                "company": company or nome,
                "job_title": g("job_title", "Cargo", "cargo"),
                "area_of_expertise_id": area_id,
                "user_id": user_id,
            }
            # validar campos mínimos
            if not data["area_of_expertise_id"] and default_area_id:
                data["area_of_expertise_id"] = default_area_id
            # normalizar email
            if isinstance(data["email"], str):
                data["email"] = data["email"].strip()
            # validar email
            # e-mail agora é opcional: validar somente se existir
            if data["email"]:
                try:
                    validate_email(str(data["email"]))
                except EmailNotValidError:
                    data["email"] = None

            # requisitos mínimos: cnpj e razao_social pelo novo layout
            if data["cnpj"] and data["razao_social"] and data["area_of_expertise_id"]:
                # evitar duplicados por email ou por razao_social
                existing = None
                if data["email"]:
                    existing = lead_repository.get_by_email(db, email=str(data["email"]))
                if not existing and data["razao_social"]:
                    # tentar encontrar pelo nome exato
                    q = db.query(Lead).filter(Lead.razao_social.ilike(str(data["razao_social"])) )
                    existing = q.first()
                if existing:
                    # não permitir troca de área automaticamente; apenas atualizar demais campos
                    update_fields = {k: v for k, v in data.items() if k not in ("user_id", "area_of_expertise_id") and v is not None}
                    for k, v in update_fields.items():
                        setattr(existing, k, v)
                    db.add(existing)
                    db.commit()
                    db.refresh(existing)
                    updated += 1
                    # progresso
                    if progress_total:
                        progress = int(((progress_base + idx) / progress_total) * 100)
                    else:
                        progress = int((idx / total) * 100)
                    from rq import get_current_job
                    job = get_current_job()
                    if job:
                        job.meta.update({"progress": progress, "inserted": created, "updated": updated, "skipped": skipped})
                        job.save_meta()
                    continue
                try:
                    lead_repository.create(db, obj_in=data)
                    created += 1
                except Exception:
                    db.rollback()
                    # upsert: atualizar registro existente
                    existing = lead_repository.get_by_email(db, email=str(data["email"])) if data["email"] else None
                    if existing:
                        update_fields = {k: v for k, v in data.items() if k not in ("user_id", "area_of_expertise_id") and v is not None}
                        for k, v in update_fields.items():
                            setattr(existing, k, v)
                        db.add(existing)
                        db.commit()
                        db.refresh(existing)
                        updated += 1
                        progress = int((idx / total) * 100)
                        from rq import get_current_job
                        job = get_current_job()
                        if job:
                            job.meta.update({"progress": progress, "inserted": created, "updated": updated, "skipped": skipped})
                            job.save_meta()
                        continue
            # progresso regular
            if progress_total:
                progress = int(((progress_base + idx) / progress_total) * 100)
            else:
                progress = int((idx / total) * 100) if total > 0 else 100
            from rq import get_current_job
            job = get_current_job()
            if job:
                job.meta.update({"progress": progress, "inserted": created, "updated": updated, "skipped": skipped})
                job.save_meta()
        # Finaliza meta com 100%
        try:
            from rq import get_current_job
            job = get_current_job()
            if job:
                job.meta.update({"progress": 100 if not progress_total else int(((progress_base + total) / progress_total) * 100), "inserted": created, "updated": updated, "skipped": skipped})
                job.save_meta()
        except Exception:
            pass
        return created
    finally:
        db.close()


def import_leads_from_file(path: str, user_id: str, default_area_id: str | None = None) -> int:
    """Importa leads processando arquivos em streaming para evitar alto uso de memória.

    Suporta CSV, XLSX e XLS. Para XLSX tenta priorizar a planilha 'Estabelecimentos'.
    Mantém atualização de progresso em job.meta e realiza upsert similar a import_leads_from_rows.
    """
    ext = os.path.splitext(path)[1].lower()

    # Inicializa meta do job
    try:
        from rq import get_current_job
        job = get_current_job()
        if job:
            job.meta.update({"progress": 0, "inserted": 0, "updated": 0, "skipped": 0})
            job.save_meta()
    except Exception:
        job = None  # type: ignore

    # Estimar total de linhas de dados (ignora cabeçalho)
    total = 0
    try:
        if ext == ".csv":
            with open(path, "rb") as f:
                total = max(0, sum(1 for _ in f) - 1)
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = None
            # prioriza planilha Estabelecimentos
            if "Estabelecimentos" in wb.sheetnames:
                ws = wb["Estabelecimentos"]
            else:
                ws = wb.active
            total = sum(1 for _ in ws.iter_rows(min_row=2, values_only=True))
            wb.close()
        elif ext == ".xls":
            book = xlrd.open_workbook(path)
            sheet = book.sheet_by_index(0)
            total = max(0, sheet.nrows - 1)
    except Exception:
        total = 0

    # Sessão de banco
    db: Session = SessionLocal()
    created = 0
    updated = 0
    skipped = 0

    def update_progress(current_index: int) -> None:
        try:
            from rq import get_current_job
            j = get_current_job()
            if j:
                progress = int((current_index / total) * 100) if total > 0 else 0
                j.meta.update({"progress": progress, "inserted": created, "updated": updated, "skipped": skipped})
                j.save_meta()
        except Exception:
            pass

    # Função utilitária para obter valores por chaves alternativas
    def g(row: dict, *keys):
        for k in keys:
            v = row.get(k)
            if v not in (None, ""):
                return v
        return None

    # Normalizador de telefone
    def norm_phone(v):
        if v in (None, ""):
            return None
        digits = ''.join([c for c in str(v) if c.isdigit()])
        if not digits:
            return None
        if digits.startswith('55') and len(digits) in (12, 13):
            return "+" + digits
        return "+55" + digits

    # Processador de uma linha em dict
    def process_row(row: dict) -> None:
        nonlocal created, updated, skipped
        # utilitário para converter valores para date
        def parse_date_value(v):
            if v in (None, ""):
                return None
            if isinstance(v, date):
                return v
            if isinstance(v, datetime):
                return v.date()
            if isinstance(v, str):
                s = v.strip()
                for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%Y/%m/%d"):
                    try:
                        return datetime.strptime(s, fmt).date()
                    except Exception:
                        pass
                return None
            return None
        # Permitir que a área seja nome (string) além de UUID
        area_value = g(row, "area_of_expertise_id", "area", "AreaId", "Area")
        if area_value and isinstance(area_value, str) and len(area_value) > 20 and '-' in area_value:
            area_id = area_value
        else:
            area = None
            if isinstance(area_value, str) and area_value.strip():
                area = area_of_expertise_repository.get_by_name(db, name=str(area_value).strip())
            area_id = str(area.id) if area else area_value

        raw_phone = g(row, "phone", "Telefone", "telefone", "Phone", "Telefone Principal", "telefone principal")
        raw_phone2 = g(row, "Telefone Secundário", "telefone secundario", "telefone secund  e1rio")
        phone_norm = norm_phone(raw_phone) or norm_phone(raw_phone2)

        company = g(row, "company", "Empresa", "empresa", "Estabelecimentos", "estabelecimentos")
        nome = g(row, "full_name", "name", "nome", "Nome", "Razão Social", "razao social", "Raz e3o Social")

        cnpj = g(row, "CNPJ", "cnpj")
        razao_social = g(row, "Raz e3o Social", "Razao Social", "razao social", "RAZAO SOCIAL")
        if not razao_social:
            razao_social = nome

        data = {
            "cnpj": cnpj,
            "razao_social": razao_social,
            "porte_codigo": g(row, "Porte C f3digo", "Porte Codigo", "Porte C f3digo"),
            "porte": g(row, "Porte"),
            "capital_social": g(row, "Capital Social"),
            "natureza_juridica_codigo": g(row, "Natureza Jur eddica C f3digo", "Natureza Juridica Codigo"),
            "natureza_juridica": g(row, "Natureza Jur eddica", "Natureza Juridica"),
            "ente_federativo_responsavel": g(row, "Ente Federativo Respons elvel", "Ente Federativo Responsavel"),
            "tipo": g(row, "Tipo"),
            "data_abertura": parse_date_value(g(row, "Data de Abertura")),
            "nome_fantasia": g(row, "Nome Fantasia"),
            "situacao_cadastral_codigo": g(row, "Situa e7 e3o Cadastral C f3digo", "Situacao Cadastral Codigo"),
            "situacao_cadastral": g(row, "Situa e7 e3o Cadastral", "Situacao Cadastral"),
            "situacao_cadastral_data": parse_date_value(g(row, "Situa e7 e3o Cadastral Data", "Situacao Cadastral Data")),
            "situacao_motivo_codigo": g(row, "Situa e7 e3o Motivo C f3digo", "Situacao Motivo Codigo"),
            "situacao_motivo": g(row, "Situa e7 e3o Motivo", "Situacao Motivo"),
            "situacao_especial_codigo": g(row, "Situa e7 e3o Especial C f3digo", "Situacao Especial Codigo"),
            "situacao_especial": g(row, "Situa e7 e3o Especial", "Situacao Especial"),
            "situacao_especial_data": parse_date_value(g(row, "Situa e7 e3o Especial Data", "Situacao Especial Data")),
            "telefone_principal": raw_phone,
            "telefone_secundario": raw_phone2,
            "phone": phone_norm,
            "email": g(row, "email", "Email", "E-mail", "e-mail"),
            "municipio_ibge": g(row, "Munic edpio IBGE", "Municipio IBGE"),
            "logradouro": g(row, "Logradouro"),
            "numero": g(row, "N famero", "Numero"),
            "complemento": g(row, "Complemento"),
            "bairro": g(row, "Bairro"),
            "cidade": g(row, "Cidade"),
            "estado": g(row, "Estado"),
            "cep": g(row, "CEP"),
            "pais_codigo": g(row, "Pais C f3digo", "Pais Codigo"),
            "pais": g(row, "Pais", "Pais"),
            "atividade_principal_codigo": g(row, "Atividade Principal C f3digo", "Atividade Principal Codigo"),
            "atividade_principal": g(row, "Atividade Principal"),
            "ultima_atualizacao": parse_date_value(g(row, " daltima Atualiza e7 e3o", "Ultima Atualizacao")),
            "full_name": nome,
            "company": company or nome,
            "job_title": g(row, "job_title", "Cargo", "cargo"),
            "area_of_expertise_id": area_id,
            "user_id": user_id,
        }

        # Completar área padrão se necessário
        if not data["area_of_expertise_id"] and default_area_id:
            data["area_of_expertise_id"] = default_area_id

        # Normalizar/validar email (opcional)
        if isinstance(data["email"], str):
            data["email"] = data["email"].strip()
        if data["email"]:
            try:
                validate_email(str(data["email"]))
            except EmailNotValidError:
                data["email"] = None

        # Requisitos mínimos
        if not (data["cnpj"] and data["razao_social"] and data["area_of_expertise_id"]):
            skipped += 1
            return

        # Upsert: priorizar CNPJ. Se não existir, criar novo.
        # Não mesclar por e-mail quando CNPJ existe e é diferente; nesse caso, evitar conflito de e-mail.
        existing = None
        if data["cnpj"]:
            existing = lead_repository.get_by_cnpj(db, cnpj=str(data["cnpj"]))
        if existing:
            update_fields = {k: v for k, v in data.items() if k not in ("user_id", "area_of_expertise_id") and v is not None}
            # evitar violação de unicidade de email ao atualizar
            if "email" in update_fields and update_fields["email"]:
                other = lead_repository.get_by_email(db, email=str(update_fields["email"]))
                if other is not None and getattr(other, "id", None) != getattr(existing, "id", None):
                    update_fields.pop("email", None)
            for k, v in update_fields.items():
                setattr(existing, k, v)
            db.add(existing)
            db.commit()
            db.refresh(existing)
            updated += 1
            return

        # Se vamos inserir e o e-mail já existe em outro lead, anular e-mail para evitar conflito com índice único parcial
        if data["email"]:
            existing_by_email = lead_repository.get_by_email(db, email=str(data["email"]))
            if existing_by_email is not None:
                data["email"] = None

        try:
            lead_repository.create(db, obj_in=data)
            created += 1
        except Exception:
            db.rollback()
            # Como último recurso: se não há CNPJ e há um lead com mesmo e-mail, atualizar esse lead
            if not data.get("cnpj") and data.get("email"):
                existing_email = lead_repository.get_by_email(db, email=str(data["email"]))
                if existing_email:
                    update_fields = {k: v for k, v in data.items() if k not in ("user_id", "area_of_expertise_id") and v is not None}
                    for k, v in update_fields.items():
                        setattr(existing_email, k, v)
                    db.add(existing_email)
                    db.commit()
                    db.refresh(existing_email)
                    updated += 1
                else:
                    skipped += 1
            else:
                skipped += 1

    try:
        processed = 0
        if ext == ".csv":
            with open(path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    processed += 1
                    process_row(row)
                    if processed % 100 == 0:
                        update_progress(processed)
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            # prioriza planilha Estabelecimentos
            if "Estabelecimentos" in wb.sheetnames:
                ws = wb["Estabelecimentos"]
            else:
                ws = wb.active

            # detectar cabeçalho
            rows_iter = ws.iter_rows(values_only=True)
            def normalize(s: str) -> str:
                if s is None:
                    return ""
                s = str(s)
                s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
                return s.strip().lower()
            header_markers = {"razao social", "nome", "email", "e-mail", "empresa", "estabelecimentos"}
            headers = None
            for row in rows_iter:
                values_norm = [normalize(c) for c in row]
                non_empty_count = sum(1 for v in values_norm if v)
                if non_empty_count >= 3 and any(v in header_markers for v in values_norm):
                    headers = [str(c).strip() if c is not None else "" for c in row]
                    break
            if headers is None:
                ws2_iter = ws.iter_rows(values_only=True)
                headers = [str(c).strip() if c is not None else "" for c in next(ws2_iter)]
                data_iter = ws2_iter
            else:
                data_iter = rows_iter

            # processa em chunks para reduzir memória/tempo de transação
            CHUNK_SIZE = 1000
            buffer = []
            for row in data_iter:
                if not any((c is not None and str(c).strip() != '') for c in row):
                    continue
                buffer.append(row)
                if len(buffer) >= CHUNK_SIZE:
                    for r in buffer:
                        processed += 1
                        row_dict = {headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))}
                        process_row(row_dict)
                        if processed % 100 == 0:
                            update_progress(processed)
                    buffer.clear()
            # restante
            if buffer:
                for r in buffer:
                    processed += 1
                    row_dict = {headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))}
                    process_row(row_dict)
                    if processed % 100 == 0:
                        update_progress(processed)
            wb.close()
        elif ext == ".xls":
            book = xlrd.open_workbook(path)
            sheet = book.sheet_by_index(0)
            def normalize(s: str) -> str:
                if s is None:
                    return ""
                s = str(s)
                s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
                return s.strip().lower()
            header_markers = {"razao social", "nome", "email", "e-mail", "empresa", "estabelecimentos"}
            header_row_index = 0
            for r in range(min(10, sheet.nrows)):
                values_norm = [normalize(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
                non_empty_count = sum(1 for v in values_norm if v)
                if non_empty_count >= 3 and any(v in header_markers for v in values_norm):
                    header_row_index = r
                    break
            headers = [str(sheet.cell_value(header_row_index, col)).strip() for col in range(sheet.ncols)]
            for r in range(header_row_index + 1, sheet.nrows):
                values = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
                if not any((c is not None and str(c).strip() != '') for c in values):
                    continue
                processed += 1
                row_dict = {headers[i]: (values[i] if i < len(values) else None) for i in range(len(headers))}
                process_row(row_dict)
                if processed % 100 == 0:
                    update_progress(processed)
        # Finaliza meta
        try:
            from rq import get_current_job
            j = get_current_job()
            if j:
                j.meta.update({"progress": 100, "inserted": created, "updated": updated, "skipped": skipped})
                j.save_meta()
        except Exception:
            pass
        return created
    finally:
        db.close()


def dispatch_campaign(campaign_id: str, lead_ids: list[str] | None = None, area_of_expertise_ids: list[str] | None = None) -> int:
    db: Session = SessionLocal()
    sent = 0
    try:
        campaign: Campaign | None = campaign_repository.get(db, campaign_id)
        if not campaign:
            return 0
        # Controle de pausa via Redis
        conn = get_redis_connection()
        paused_key = f"campaign:{campaign_id}:paused"
        if conn.get(paused_key):
            return 0
        q = db.query(Lead)
        if lead_ids:
            q = q.filter(Lead.id.in_(lead_ids))
        if area_of_expertise_ids:
            q = q.filter(Lead.area_of_expertise_id.in_(area_of_expertise_ids))
        leads = q.all()
        # Respeitar canais selecionados
        channels_list = list(getattr(campaign, 'channels', []) or [])
        channel_names = {getattr(c, 'name', '').lower() for c in channels_list}
        send_email_enabled = 'e-mail' in channel_names or 'email' in channel_names
        # Descobrir IDs do canal de e-mail (se existirem)
        email_channel_ids: list[str] = []
        try:
            for ch in channels_list:
                n = (getattr(ch, 'name', '') or '').lower()
                if n in ('e-mail', 'email'):
                    email_channel_ids.append(str(getattr(ch, 'id', '')))
        except Exception:
            email_channel_ids = []
        send_sms_enabled = 'sms' in channel_names
        base_url = os.getenv("BASE_URL", "http://localhost:8000").rstrip("/")
        for lead in leads:
            # parar se estiver pausado
            if conn.get(paused_key):
                break
            ok = False
            if send_email_enabled:
                # Pré-cria registro para obter o ID e usá-lo no pixel de rastreamento
                # Tenta mapear o canal E-mail (se existir) para o envio
                email_channel_id = None
                try:
                    for ch in channels_list:
                        n = (getattr(ch, 'name', '') or '').lower()
                        if n in ('e-mail', 'email'):
                            email_channel_id = getattr(ch, 'id', None)
                            break
                except Exception:
                    email_channel_id = None

                cl = CampaignLead(campaign_id=campaign.id, lead_id=lead.id, status=EmailStatus.PENDING)
                if email_channel_id:
                    try:
                        cl.channel_id = email_channel_id
                    except Exception:
                        pass
                db.add(cl)
                # Buscar template específico por canal (email); fallback: geral (channel_id IS NULL); fallback final: body/subject da campanha
                tpl = None
                try:
                    from src.models.template import Template
                    if email_channel_ids:
                        tpl = (
                            db.query(Template)
                            .filter(Template.campaign_id == campaign.id, Template.channel_id.in_(email_channel_ids))
                            .first()
                        )
                    if not tpl:
                        tpl = (
                            db.query(Template)
                            .filter(Template.campaign_id == campaign.id, Template.channel_id.is_(None))
                            .first()
                        )
                except Exception:
                    tpl = None
                subject = (getattr(tpl, 'subject', None) or campaign.subject or '')
                body = (getattr(tpl, 'content', None) or campaign.body or '')
                attachments = []
                try:
                    for a in getattr(tpl, 'attachments', []) or []:
                        attachments.append({
                            'file_name': getattr(a, 'file_name', None),
                            'file_url': getattr(a, 'file_url', None),
                            'content_type': getattr(a, 'content_type', None),
                        })
                except Exception:
                    attachments = []
                # Incluir pixel de rastreamento de abertura vinculado à campanha e ao envio (campaign_lead)
                try:
                    tracking_pixel = (
                        f"<img src='{base_url}/api/campaigns/{str(campaign.id)}/events/open/{str(cl.id)}' width='1' height='1' style='display:none' alt=''/>"
                    )
                except Exception:
                    tracking_pixel = ""
                body = (body or "") + tracking_pixel
                ok = False
                error_msg = None
                try:
                    # Se a campanha tiver conta de marketing, configurar credenciais SMTP
                    try:
                        if getattr(campaign, 'marketing_account_id', None):
                            acc_id = str(getattr(campaign, 'marketing_account_id'))
                            cred = marketing_service.get_smtp(db, account_id=acc_id)
                            if cred:
                                email_service.use_account(
                                    host=getattr(cred, 'host', None) or os.getenv('MAIL_HOST', 'localhost'),
                                    port=getattr(cred, 'port', None) or os.getenv('MAIL_PORT', '25'),
                                    encryption=getattr(cred, 'encryption', None) or os.getenv('MAIL_ENCRYPTION', ''),
                                    username=decrypt_str(getattr(cred, 'username_enc', None)),
                                    password=decrypt_str(getattr(cred, 'password_enc', None)),
                                    from_name=getattr(cred, 'from_name', None) or os.getenv('MAIL_FROM_NAME', None),
                                    from_address=getattr(cred, 'from_address', None) or os.getenv('MAIL_FROM_ADDRESS', None),
                                    reply_to=getattr(cred, 'reply_to', None),
                                )
                    except Exception:
                        pass
                    ok = email_service.send_email(to=lead.email, subject=subject, body=body, attachments=attachments)
                except Exception as e:
                    ok = False
                    error_msg = str(e)
                finally:
                    try:
                        email_service.reset_account()
                    except Exception:
                        pass
                if ok:
                    cl.status = EmailStatus.SENT
                    cl.sent_at = datetime.utcnow()
                    sent += 1
                else:
                    cl.status = EmailStatus.FAILED
                    # Tenta armazenar pelo menos um motivo simples
                    try:
                        if not error_msg and (not lead.email or '@' not in (lead.email or '')):
                            cl.error_reason = 'E-mail inválido ou ausente'
                        elif error_msg:
                            cl.error_reason = error_msg[:255]
                    except Exception:
                        pass
        # Atualiza status da campanha quando houver pelo menos 1 envio
        if sent > 0:
            try:
                campaign.status = CampaignStatus.SENT
                campaign.sent_at = datetime.utcnow()
                db.add(campaign)
            except Exception:
                pass
        db.commit()
        return sent
    finally:
        db.close()


